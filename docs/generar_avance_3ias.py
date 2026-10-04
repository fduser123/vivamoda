#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera Avance_Detallado_Examen_Parcial_VivaModa_3_IA.docx

Espejo estructural del documento de referencia
'Avance_Detallado_Examen_Parcial_Carwash_Opcion1.pdf', adaptado al proyecto
web de moda VivaModa (Node + Express + PostgreSQL) y a sus 3 funciones de IA:

  1) Chatbot estilista con memoria y contexto (Aria)
  2) Asesor visual por imagen (recomendación de prendas)
  3) Ayudante IA del área de administración (panel de almacén)

Requiere: python-docx  (pip install python-docx)
Uso:      python3 generar_avance_3ias.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------- paleta
BRAND = RGBColor(0xB6, 0x00, 0x55)   # primary VivaModa
PURPLE = RGBColor(0x4B, 0x41, 0xE1)   # secondary
INK = RGBColor(0x1C, 0x1B, 0x1D)
MUTED = RGBColor(0x5C, 0x3F, 0x45)

doc = Document()

# ---------------------------------------------------------------- estilos
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.font.color.rgb = INK

for lvl, (sz, color) in enumerate([(16, BRAND), (13, BRAND), (11.5, PURPLE)], start=1):
    h = doc.styles['Heading %d' % lvl]
    h.font.name = 'Calibri'
    h.font.size = Pt(sz)
    h.font.color.rgb = color
    h.font.bold = True

sec = doc.sections[0]
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)


# ---------------------------------------------------------------- helpers
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def shade(cell, hexcolor):
    """Aplica color de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(W + 'val', 'clear')
    shd.set(W + 'color', 'auto')
    shd.set(W + 'fill', hexcolor)
    tcPr.append(shd)


def P(text='', size=11, bold=False, italic=False, color=None,
      align=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or INK
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    return p


def BUL(text, size=11):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.name = 'Calibri'
    return p


def KV(label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = MUTED
    r2 = p.add_run(value)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(2)
    return p


def TABLE(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], 'B60055')
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            para = cells[i].paragraphs[0]
            run = para.add_run(val)
            run.font.size = Pt(9.5)
            run.font.name = 'Calibri'
            if r_i % 2 == 1:
                shade(cells[i], 'F6F2F5')
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def CALLOUT(title, text, fill='F3E5EE'):
    """Recuadro destacado, equivalente a las cajas del PDF de referencia."""
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    cell.text = ''
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = BRAND
    p2 = cell.add_paragraph()
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = INK
    shade(cell, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ================================================================ PORTADA
P('ESPECIFICACIÓN TÉCNICA Y GUÍA DE AVANCE DEL PROYECTO',
  size=17, bold=True, color=BRAND, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
P('Avance del Examen Parcial: Sistema Web VivaModa con Integración de IA '
  '(Opción 1 - Enfoque Ligero API)',
  size=13, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

P('Curso: Gobiernos y Gestión de TI / Inteligencia Artificial\t\t'
  'Evaluación: Avance de Examen Parcial',
  size=10, color=MUTED, space_after=2)
P('Arquitectura Target: Opción 1 - API Externa (Node.js + Express + '
  'PostgreSQL + LLM API)\t\tFecha: Semestre Académico 2026-II',
  size=10, color=MUTED, space_after=2)
P('Módulos de IA: 1) Chatbot Aria  2) Asesor Visual por Imagen  '
  '3) Ayudante IA de Administración',
  size=10, color=MUTED, space_after=12)

# ================================================================ 1
doc.add_heading('1. Definición del Problema Complejo', level=1)
P('El comercio de moda omnicanal (tienda web, punto de venta físico, almacén '
  'y venta por catálogo) enfrenta ineficiencias dinámicas en la experiencia de '
  'compra y en la operación administrativa. Los modelos tradicionales de e-commerce '
  'sufren de asesoría impersonal, compras mal informadas sobre qué prendas favorecen '
  'a cada cuerpo y decisiones de compra o reposición tomadas a ciegas.')
P('A continuación se detallan los tres problemas complejos fundamentales que el '
  'sistema resuelve:')
BUL('Asesoría Impersonal y Recomendaciones Genéricas: el cliente no recibe sugerencias '
    'basadas en su historial de compras, tallas, preferencias declaradas ni en el '
    'catálogo disponible, lo que desperdicia oportunidades de venta cruzada y de '
    'valorización de prendas: el cliente debe deducir por su cuenta qué le favorece.')
BUL('Decisiones de Inventario a Cegas: la administración de almacén no cuenta con previsión '
    'de demanda ni con detección de anomalías, por lo que acumula quiebres de stock '
    'en productos clave y sobredimensiona artículos de lenta rotación (SKU estancados).')
BUL('Carga Operativa Repetitiva en Soporte y Administración: las consultas '
    'frecuentes (precios, stock, estado de pedidos, reportes) y los ajustes manuales '
    'de inventario consumen tiempo del personal, restando a la capacidad de atención '
    'comercial del equipo.')

# ================================================================ 2
doc.add_heading('2. Definición del Aplicativo Web Básico', level=1)
P('El proyecto se estructura como una aplicación web monolítica ligera pero escalable, '
  'diseñada para operar de forma ágil bajo la pila de tecnologías estándar:')
TABLE(
    ['Capa de Arquitectura', 'Tecnología Seleccionada', 'Rol en el Sistema'],
    [
        ['Frontend (Interfaz)',
         'HTML5, Tailwind CSS, JavaScript (Vanilla)',
         'Interfaz dinámica para clientes, personal de tienda y administradores. '
         'Captura de eventos y actualización asíncrona (Fetch API + SSE).'],
        ['Backend (Lógica)',
         'Node.js + Express (servidor HTTP)',
         'Procesamiento de lógica de negocio, autenticación por roles, orquestación '
         'de consultas y comunicación HTTPS con la API de IA.'],
        ['Base de Datos',
         'PostgreSQL',
         'Almacenamiento relacional de clientes, productos, variantes, inventario, '
         'pedidos, carrito y memoria/historial de las interacciones con la IA.'],
    ],
    widths=[1.5, 2.0, 3.2])

# ================================================================ 3
doc.add_heading('3. Caracterización de Requerimientos Funcionales con IA (3 Funciones)', level=1)
P('En concordancia con los lineamientos del examen parcial, la Opción 1 integra '
  'Inteligencia Artificial Generativa y Predictiva a través de API externa en tres '
  'módulos del sistema:')

doc.add_heading('RF-01: Chatbot Estilista con Memoria y Contexto (Aria)', level=2)
BUL('Descripción: Asistente conversacional que atiende dudas de catálogo, tallas, '
    'colores y estado de pedidos, y que además aprende de las preferencias declaradas '
    'por el cliente (afirmaciones y negaciones sobre prendas, colores y estilos).')
BUL('Subproceso / Actividades: El cliente escribe un mensaje en el widget de chat. '
    'JavaScript reenvía la consulta al backend, el Server recupera la memoria '
    'persistida del cliente desde PostgreSQL, consolida el contexto (perfil, pedidos, '
    'carrito, catálogo y resultados de búsqueda web) y lo inyecta en el prompt antes '
    'de llamar a la API de IA. La respuesta se transmite en streaming al navegador.')
BUL('Componente IA: Modelo de lenguaje conversacional con System Prompt delimitado '
    '(guardrails de contexto) y extracción estructurada de hechos (memoria).')

doc.add_heading('RF-02: Asesor Visual por Imagen', level=2)
BUL('Descripción: Analiza la fotografía de una prenda cargada por el usuario y devuelve '
    'una descripción estructurada (tipo de prenda, color dominante, patrón,occasion) '
    'que alimenta el ranking del catálogo y las sugerencias de estilo.')
BUL('Subproceso / Actividades: El frontend valida la imagen y la envía como payload. '
    'El backend la procesa (descripción y clasificación visual) y, junto con el '
    'catálogo activo y las preferencias del usuario, genera recomendaciones '
    'personalizadas que se renderizan en una tarjeta de estilo dinámico.')
BUL('Componente IA: Integración de API de IA con salida estructurada (JSON Mode) para '
    'descripción visual y recomendación de productos afines.')

doc.add_heading('RF-03: Ayudante IA del Área de Administración', level=2)
BUL('Descripción: Módulo de asistencia al panel de almacén y ventas que provee previsión '
    'de demanda, detección de anomalías de inventario, clasificación ABC de productos '
    'y una estrategia de compras priorizada, además de responder preguntas en lenguaje '
    'natural sobre los datos del negocio.')
BUL('Subproceso / Actividades: El administrador abre el panel autenticado. El backend '
    'calcula los indicadores (pronóstico de quiebres, anomalías, ABC por ingresos, SKU '
    'estancados) sobre PostgreSQL, los entrega al frontend y este los muestra en '
    'pestañas con KPIs y plan de compra accionable.')
BUL('Componente IA: Modelos de lenguaje con analítica aumentada (text-to-SQL de solo '
    'lectura con enmascaramiento de datos personales de clientes).')

# ================================================================ 4
doc.add_heading('4. Diagrama de Base de Datos Relacional (PostgreSQL)', level=1)
P('El esquema relacional en PostgreSQL soporta la operación básica y almacena las trazas '
  'e interacciones generadas por la Inteligencia Artificial.')
TABLE(
    ['Tabla', 'Campos Clave / Atributos', 'Relación / Descripción'],
    [
        ['users', 'id, full_name, email, password_hash, role, vip_tier, interests, points',
         '1 a N con orders. Roles: customer / staff / admin. Guarda preferencias declaradas.'],
        ['products', 'id, sku, name, gender, category, price, image_url, visibility, badge',
         'Catálogo de prendas. 1 a N con product_variants.'],
        ['product_variants / inventory', 'id, product_id, size, color, qty, reorder_point',
         'Variantes y existencias por tienda. Alimenta el panel de almacén y la tienda VR.'],
        ['orders / order_items', 'id, user_id, store_id, status, total, created_at; sku, qty, price',
         'Historial transaccional. Base del análisis de compra para las 3 IAs.'],
        ['carts / cart_items', 'id, user_id; sku, size, qty',
         'Carrito activo: contexto inmediato para el chatbot y recomendaciones.'],
        ['ai_sessions / ai_events', 'id, session_key, role, kind, payload, created_at',
         'Trazas de las interacciones con la IA. 1 a N con users.'],
        ['ai_memory', 'id, user_id, kind, fact_key, fact_value, confidence, source, updated_at',
         'Memoria persistente que la IA aprende del cliente (hechos y preferencias).'],
        ['ai_web_cache', 'key, payload, fetched_at',
         'Caché de resultados de búsqueda web usados para enriquecer el contexto.'],
    ],
    widths=[1.4, 2.4, 2.9])

# ================================================================ 5
doc.add_heading('5. Procedimiento Detallado de Integración (API Externa)', level=1)
P('El flujo de integración técnica bajo el Enfoque Ligero de la Opción 1 se divide en '
  '5 etapas secuenciales:')
BUL('Paso 1 — Configuración de Credenciales y Entorno: Se genera una API Key en el '
    'proveedor (Google Gemini / OpenRouter / OpenAI). Se almacena dentro del archivo '
    'de configuración del servidor backend (.env) garantizando que nunca se exponga '
    'en el código JavaScript del cliente.')
BUL('Paso 2 — Extracción de Contexto Local (SQL): Al activarse un evento en la web, el '
    'backend ejecuta consultas a PostgreSQL para extraer el perfil del cliente, su '
    'historial de pedidos, el contenido del carrito, el inventario y la memoria '
    'aprendida.')
BUL('Paso 3 — Construcción y Delimitación del Prompt (System Prompting): El backend '
    'construye una estructura de prompt que incluye: (a) Rol de la IA, (b) Reglas '
    'operativas estrictas y delimitadores de contexto, (c) Datos del cliente extraídos '
    'de PostgreSQL, y (d) Formato de salida requerido.')
BUL('Paso 4 — Consumo REST/HTTPS con fetch: El servidor realiza una solicitud POST '
    'asíncrona hacia la API Externa enviando el payload JSON codificado con '
    'autenticación Bearer Token (nunca desde el navegador).')
BUL('Paso 5 — Validación, Almacenamiento y Renderizado: El backend recibe la respuesta, '
    'valida la estructura, guarda la traza y la memoria aprendida en PostgreSQL y '
    'transmite los datos procesados al Frontend para su despliegue (incluido streaming).')

# ================================================================ 6
doc.add_heading('6. Explicación del Funcionamiento Arquitectónico (Backend y Frontend)', level=1)

doc.add_heading('Flujo Operativo en el Backend (Node + PostgreSQL)', level=2)
BUL('Custodia de Seguridad: El Backend actúa como intermediario seguro (Proxy), '
    'evitando que las API Keys de la IA se expongan en el navegador del usuario.')
BUL('Consolidación de Contexto: Transforma los datos relacionales planos de PostgreSQL '
    'en objetos estructurados aptos para consumo de LLMs.')
BUL('Control de Excepciones: Si la API de IA experimenta latencia o caída (Timeout), el '
    'backend captura la excepción y retorna una respuesta por defecto (Fallback) '
    'garantizando que el sistema VivaModa continúe funcionando sin interrupciones.')

doc.add_heading('Flujo Operativo en el Frontend (HTML / JS / CSS)', level=2)
BUL('Eventos de Usuario: Captura acciones de la interfaz mediante JavaScript (ej. enviar '
    'un mensaje al chat, cambiar el selector de productos, cargar una imagen).')
BUL('Consultas Asíncronas (AJAX / Fetch API): Envía peticiones HTTP no bloqueantes hacia '
    'las rutas expuestas por el Backend.')
BUL('Renderizado Dinámico: Manipula el DOM para mostrar tarjetas de recomendación, '
    'indicadores de stock en el panel, burbujas en el chat de soporte y modales de '
    'confirmación.')

# ================================================================ REQUISITO
CALLOUT(
    'REQUISITO EVALUATIVO PARA EL EXAMEN PARCIAL',
    'Para la entrega del Examen Parcial, los alumnos deberán presentar la ejecución '
    'funcional de estos 6 puntos en el aplicativo web, demostrando que al menos las '
    '3 funciones de IA consumen la API externa desde el Backend y reflejan los datos '
    'dinámicamente en la interfaz.',
    fill='F3E5EE')

CALLOUT(
    'Demostración funcional — las 3 funciones de IA',
    'RF-01 Chatbot Aria: ruta /api/ai/chat/stream (SSE) y /api/ai/chat; refleja burbujas '
    'de chat, recomendaciones de catálogo y memoria aprendida en el panel.\n'
    'RF-02 Asesor Visual: carga de imagen en el detalle de producto; devuelve descripción '
    'y ranking de prendas afines renderizado como tarjeta de estilo.\n'
    'RF-03 Ayudante de Administración: rutas /api/admin/ai/insights y /api/admin/ai/memory; '
    'refleja previsión de demanda, anomalías, clasificación ABC y plan de compras en el '
    'panel de almacén.',
    fill='EDE7F6')

# ================================================================ PERSISTENCIA
doc.add_heading('7. Persistencia de las Interacciones con la IA', level=1)
P('Cada interacción con la IA queda registrada para alimentar la memoria del sistema y '
  'garantizar la trazabilidad de las recomendaciones:')
BUL('ai_sessions: agrupa los mensajes por cliente y sesión, permitiendo continuidad del '
    'contexto entre recargas.')
BUL('ai_events: registra eventos de interacción con su carga útil (payload) para análisis.')
BUL('ai_memory: almacena los hechos aprendidos (preferencias, tallas, categorías, colores) '
    'con su nivel de confianza y origen, y es editable por el administrador.')
BUL('ai_web_cache: cachea temporalmente los resultados de búsqueda web que enriquecen el '
    'contexto del chatbot.')

# ================================================================ cierre
doc.add_heading('8. Conclusiones', level=1)
P('El proyecto VivaModa integra tres módulos de Inteligencia Artificial '
  '(chatbot estilista con memoria, asesor visual por imagen y ayudante de administración) '
  'mediante un backend proxy seguro que protege las credenciales, consolida el contexto '
  'desde PostgreSQL y degrada de forma elegante ante fallos del proveedor de IA. '
  'El frontend captura eventos de usuario, consume las rutas del backend mediante '
  'peticiones asíncronas y refleja los resultados de la IA de forma dinámica en la '
  'interfaz, cumpliendo el requisito de ejecución funcional de los 6 puntos.')

# ---------------------------------------------------------------- guardar
OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'Avance_Detallado_Examen_Parcial_VivaModa_3_IA.docx')
doc.save(OUT)
print('OK: generado %s (%d bytes)' % (OUT, os.path.getsize(OUT)))