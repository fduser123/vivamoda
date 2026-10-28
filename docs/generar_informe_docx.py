#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera la versión Word del informe de revisión de IA de VivaModa.

Lee `INFORME-PRACTICA-CALIFICADA-II-VIVAMODA.md` y produce
`INFORME-PRACTICA-CALIFICADA-II-VIVAMODA.docx` con portada, índice,
tablas, bloques de código y el mismo relato del Markdown.

Requiere: python-docx (docs/.venv-doc).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
MD = AQUI / "INFORME-PRACTICA-CALIFICADA-II-VIVAMODA.md"
OUT = AQUI / "INFORME-PRACTICA-CALIFICADA-II-VIVAMODA.docx"

BRAND = RGBColor(0xB6, 0x00, 0x55)
PURPLE = RGBColor(0x4B, 0x41, 0xE1)
INK = RGBColor(0x1C, 0x1B, 0x1D)
MUTED = RGBColor(0x6B, 0x5A, 0x62)
LIGHT = RGBColor(0x8A, 0x7F, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_INK = RGBColor(0x6D, 0x1B, 0x45)
LINK = RGBColor(0x11, 0x55, 0xCC)

BODY_SIZE = 10.5
TABLE_WIDTH = 16.6  # cm útiles con márgenes de 2,2 cm en A4
LANG = {"yaml", "bash", "html", "sql", "env", "gitignore", "python", "json", "text", "mermaid"}

CAPTIONS = {
    "CAPÍTULO I": "Problemática, justificación y objetivos",
    "CAPÍTULO II": "Requerimientos funcionales inteligentes",
    "CAPÍTULO III": "Arquitectura técnica multinivel y matriz de 10 capas",
    "CAPÍTULO IV": "Generador de API y módulos embebibles",
    "CAPÍTULO V": "Asistentes, avatar 3D y persistencia híbrida",
    "CAPÍTULO VI": "MLOps, validación y plan de pruebas",
    "CAPÍTULO VII": "Conclusiones, recomendaciones y anexos",
    "ANEXO A": "Revisión técnica del subsistema de IA (18 hallazgos)",
    "ANEXO B": "Plan de evolución hacia deep learning",
    "ANEXO C": "Cumplimiento de la plantilla y del balotario",
    "ANEXO D": "Diccionario de métricas y tablero",
    "ANEXO E": "Referencias y trazabilidad",
    "ANEXO F": "Lista priorizada de mejoras implementables",
}


# ------------------------------------------------------------------
# Helpers XML
# ------------------------------------------------------------------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
        "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
    )


def cell_margins(table, top=55, start=85, bottom=55, end=85):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")


def table_borders(table, color="E4DDE3", size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"
    )


def cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.insert_element_before(
            borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
            "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
        )
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "E4DDE3"))
        borders.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def bottom_border(paragraph, color="E4DDE3", size=6, space=2):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    pPr.insert_element_before(
        bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def add_field(paragraph, instruction, placeholder="1", size=8, color=LIGHT):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = color
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)


def add_hyperlink(paragraph, url, text, size=10.5, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    if bold:
        rPr.append(OxmlElement("w:b"))
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def run_shade(run, fill="F4EFF3"):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


# ------------------------------------------------------------------
# Inline Markdown
# ------------------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+`|\[[^\]\n]+\]\([^\)\n]+\))")


def add_inline(paragraph, text, size=BODY_SIZE, bold=False, color=INK, code_size=None):
    code_size = code_size or max(7.2, size - 1.3)
    text = text.replace("\\|", "|").replace("\\*", "*").replace("\\_", "_")
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color
        elif token.startswith("`") and token.endswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(code_size)
            r.font.color.rgb = CODE_INK
            run_shade(r)
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            add_hyperlink(paragraph, url, label, size=size, bold=bold)
        else:
            r = paragraph.add_run(token[1:-1])
            r.font.size = Pt(size)
            r.font.italic = True
            r.font.bold = bold
            r.font.color.rgb = color
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color


def para(text="", size=BODY_SIZE, bold=False, color=INK, align=None,
         before=0, after=4, indent=None, first=None, keep=False):
    p = doc.add_paragraph()
    if text:
        add_inline(p, text, size=size, bold=bold, color=color)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first is not None:
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep
    return p


# ------------------------------------------------------------------
# Documento
# ------------------------------------------------------------------
doc = Document()

# Propiedades
doc.core_properties.title = "VivaModa · Informe de revisión y plan de evolución del subsistema de IA"
doc.core_properties.subject = "Práctica Calificada II · Arquitectura de 10 capas · Deep Learning"
doc.core_properties.author = "VivaModa"
doc.core_properties.keywords = "IA, deep learning, RAG, pgvector, MLOps, moda, VivaModa"

# Página A4 y márgenes
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)
sec.different_first_page_header_footer = True

# Estilos base
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(BODY_SIZE)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for name, size, color, before, after in (
    ("Heading 1", 17, BRAND, 0, 8),
    ("Heading 2", 13, BRAND, 12, 4),
    ("Heading 3", 11, PURPLE, 9, 3),
):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


def heading(text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    if level == 1:
        p.paragraph_format.page_break_before = True
        bottom_border(p, color="E8DDE4", size=6, space=4)
    elif level == 2:
        bottom_border(p, color="F0E8ED", size=3, space=2)
    return p


# Encabezado y pie (no en portada)
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("VivaModa · Informe de revisión del subsistema de IA")
hr.font.size = Pt(8)
hr.font.color.rgb = LIGHT
bottom_border(header, color="E8DDE4", size=3, space=2)

footer = sec.footer.paragraphs[0]
footer.paragraph_format.tab_stops.add_tab_stop(Cm(TABLE_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
fr = footer.add_run("Práctica Calificada II · 2026")
fr.font.size = Pt(8)
fr.font.color.rgb = LIGHT
footer.add_run("\t")
pr = footer.add_run("Página ")
pr.font.size = Pt(8)
pr.font.color.rgb = LIGHT
add_field(footer, "PAGE")
nr = footer.add_run(" de ")
nr.font.size = Pt(8)
nr.font.color.rgb = LIGHT
add_field(footer, "NUMPAGES", placeholder="1")


# ------------------------------------------------------------------
# Portada
# ------------------------------------------------------------------
def cover_band(text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.text = ""
    shade(cell, "B60055")
    cell_margins(t, top=60, start=100, bottom=60, end=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return t


cover_band("VIVAMODA  ·  PRÁCTICA CALIFICADA II")
para("", after=10)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(28)
r = p.add_run("Informe de revisión\ny plan de evolución\ndel subsistema de IA")
r.font.size = Pt(26)
r.font.bold = True
r.font.color.rgb = INK
p.paragraph_format.line_spacing = 1.02
p.paragraph_format.space_after = Pt(12)

p = doc.add_paragraph()
r = p.add_run("Aplicación Web Inteligente con Deep Learning\ny Arquitectura de 10 Capas")
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = BRAND
p.paragraph_format.space_after = Pt(10)

rule = doc.add_paragraph()
rule.paragraph_format.space_after = Pt(18)
bottom_border(rule, color="B60055", size=12, space=1)

# Metadatos de la portada
meta = {
    "Proyecto": "VivaModa · Ecosistema omnicanal de moda",
    "Asignatura": "Inteligencia Artificial — II (Mitad de ciclo)",
    "Entregable": "Informe de revisión técnica + Plan de Proyecto Final",
    "Versión": "1.0",
    "Fecha": "27 de octubre de 2026",
    "Repositorio": "proyecto 1 · rama feat/ia-multimodal-paginacion-v2",
    "Autores": "[Nombre completo] · [Código] · [Correo institucional]",
    "Sede": "[Ciudad], [País]",
    "Documento guía": "Plantilla del Plan_Proyecto II_Arquitectura_10_Capas_IA.pdf",
}
meta_table = doc.add_table(rows=0, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
meta_table.autofit = False
table_borders(meta_table, color="EFE7EC", size=3)
cell_margins(meta_table, top=70, start=0, bottom=70, end=80)
for key, value in meta.items():
    row = meta_table.add_row().cells
    row[0].width = Cm(4.1)
    row[1].width = Cm(12.4)
    p0 = row[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(key)
    r0.font.size = Pt(8.5)
    r0.font.bold = True
    r0.font.color.rgb = MUTED
    p1 = row[1].paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(value)
    r1.font.size = Pt(9.2)
    r1.font.color.rgb = INK

para("", after=18)
note = para(
    "Documento de trabajo elaborado a partir de la plantilla del Plan Proyecto II. "
    "Las cifras de rendimiento provienen de los informes del repositorio (07-10-2026) "
    "y no se re-ejecutaron en esta revisión.",
    size=8.6, color=MUTED, after=2,
)
para("VivaModa · 2026", size=9, bold=True, color=BRAND, after=0)


# ------------------------------------------------------------------
# Índice
# ------------------------------------------------------------------
doc.add_page_break()
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Contenido")
r.font.size = Pt(22)
r.font.bold = True
r.font.color.rgb = BRAND
bottom_border(p, color="B60055", size=8, space=4)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
add_inline(
    intro,
    "El informe sigue el índice oficial de la plantilla (capítulos I a VII) y añade seis anexos "
    "con la revisión técnica, el plan de deep learning y el backlog priorizado.",
    size=9.5, color=MUTED,
)

toc = doc.add_table(rows=1, cols=2)
toc.alignment = WD_TABLE_ALIGNMENT.CENTER
toc.autofit = False
table_borders(toc, color="EFE7EC", size=3)
cell_margins(toc, top=60, start=90, bottom=60, end=90)
hdr = toc.rows[0].cells
hdr[0].width = Cm(5.2)
hdr[1].width = Cm(11.3)
for cell, text in zip(hdr, ("Sección", "Qué encontrarás")):
    cell.text = ""
    shade(cell, "B60055")
    r = cell.paragraphs[0].add_run(text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = WHITE
for key, desc in CAPTIONS.items():
    cells = toc.add_row().cells
    cells[0].width = Cm(5.2)
    cells[1].width = Cm(11.3)
    p0 = cells[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    label = key.split(" ", 1)[0].capitalize() + " " + key.split(" ", 1)[1]
    r0 = p0.add_run(label)
    r0.font.size = Pt(8.6)
    r0.font.bold = True
    r0.font.color.rgb = BRAND
    p1 = cells[1].paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(desc)
    r1.font.size = Pt(8.8)
    r1.font.color.rgb = INK

para("", after=4)
para(
    "Los capítulos y anexos empiezan en página nueva. Las tablas repiten su cabecera "
    "cuando continúan en la página siguiente.",
    size=8.4, color=MUTED, after=0,
)
doc.add_page_break()


# ------------------------------------------------------------------
# Parseo del Markdown
# ------------------------------------------------------------------
lines = MD.read_text(encoding="utf-8").splitlines()
start = next(i for i, l in enumerate(lines) if l.startswith("## Cómo leer este informe"))

TABLE_SEP = re.compile(r"^:?-{2,}:?$")


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", s)]


def is_separator(line):
    if not line.strip().startswith("|"):
        return False
    cells = split_row(line)
    return len(cells) >= 2 and all(TABLE_SEP.match(c.replace(" ", "")) for c in cells)


def add_table(header, rows):
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)
    table_borders(table, color="E4DDE3", size=3)
    cell_margins(table, top=45, start=65, bottom=45, end=65)
    repeat_header(table.rows[0])

    size = 7.4 if cols >= 7 else 8.1 if cols >= 5 else 8.8

    # Ancho mínimo por palabra: evita que la cabecera se parta
    # ("Severid/ad", "Esfuerz/o", "Horizont/e").
    char_cm = size * 0.021
    mins, weights = [], []
    for j in range(cols):
        values = [header[j]] + [r[j] if j < len(r) else "" for r in rows]
        longest_word = max(
            (len(w) for v in values for w in re.findall(r"[^\s]+", v)),
            default=1,
        )
        mins.append(min(2.9, max(1.0, longest_word * char_cm + 0.28)))
        weights.append(max(8, min(60, max(len(v) for v in values))))
    if sum(mins) >= TABLE_WIDTH:
        widths = [m * TABLE_WIDTH / sum(mins) for m in mins]
    else:
        extra = TABLE_WIDTH - sum(mins)
        widths = [mins[j] + extra * weights[j] / sum(weights) for j in range(cols)]

    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)
    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.width = Cm(widths[j])
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, "B60055")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, size=size, bold=True, color=WHITE, code_size=size - 0.7)

    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(cols):
            cell = cells[j]
            cell.width = Cm(widths[j])
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if idx % 2 == 1:
                shade(cell, "FBF8FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, row[j] if j < len(row) else "", size=size, color=INK, code_size=size - 0.7)
    para("", after=3)


def add_code_block(code_lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fixed_layout(t)
    table_borders(t, color="E8E2E6", size=3)
    cell_margins(t, top=70, start=100, bottom=70, end=100)
    t.columns[0].width = Cm(TABLE_WIDTH)
    cell = t.rows[0].cells[0]
    cell.width = Cm(TABLE_WIDTH)
    cell.text = ""
    shade(cell, "F7F6F8")
    first = True
    for line in code_lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(7.2)
        r.font.color.rgb = RGBColor(0x2B, 0x25, 0x2A)
    para("", after=3)


def add_callout(text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fixed_layout(t)
    cell_margins(t, top=90, start=130, bottom=90, end=130)
    t.columns[0].width = Cm(TABLE_WIDTH)
    cell = t.rows[0].cells[0]
    cell.width = Cm(TABLE_WIDTH)
    cell.text = ""
    shade(cell, "F6F2F5")
    cell_borders(cell, left={"sz": 24, "color": "B60055"}, top={"sz": 2, "color": "F6F2F5"},
                 bottom={"sz": 2, "color": "F6F2F5"}, right={"sz": 2, "color": "F6F2F5"})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    add_inline(p, text, size=9.2, color=INK)
    para("", after=3)


def add_rule():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    bottom_border(p, color="E8DDE4", size=3, space=1)


def is_special(s):
    return (
        not s
        or s.startswith("#")
        or s.startswith("|")
        or s.startswith("```")
        or s.startswith(">")
        or s == "---"
        or s.startswith("- ")
        or re.match(r"^\d+\.\s", s)
    )


i = start
while i < len(lines):
    raw = lines[i]
    s = raw.strip()

    if not s:
        i += 1
        continue

    # Código
    if s.startswith("```"):
        i += 1
        code = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code.append(lines[i].rstrip())
            i += 1
        i += 1
        if code and code[0].strip().lower() in LANG:
            code = code[1:]
        add_code_block(code)
        continue

    # Tabla
    if s.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
        header = split_row(lines[i])
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(split_row(lines[i]))
            i += 1
        add_table(header, rows)
        continue

    # Encabezados
    m = re.match(r"^(#{1,3})\s+(.*)$", s)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        if level == 1:
            heading(text, 1)
        elif level == 2:
            heading(text, 2)
        else:
            heading(text, 3)
        i += 1
        continue

    # Regla
    if s == "---":
        add_rule()
        i += 1
        continue

    # Cita / callout
    if s.startswith(">"):
        quote = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            quote.append(lines[i].strip().lstrip(">").strip())
            i += 1
        add_callout(" ".join(quote))
        continue

    # Viñetas
    if s.startswith("- "):
        while i < len(lines) and lines[i].strip().startswith("- "):
            text = lines[i].strip()[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run("•  ")
            r.font.size = Pt(BODY_SIZE)
            r.font.color.rgb = BRAND
            add_inline(p, text, size=BODY_SIZE)
            i += 1
        continue

    # Listas numeradas: se conserva el número visible
    if re.match(r"^\d+\.\s", s):
        while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
            text = lines[i].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            add_inline(p, text, size=BODY_SIZE)
            i += 1
        continue

    # Párrafo: unir líneas consecutivas
    buf = [s]
    i += 1
    while i < len(lines):
        nxt = lines[i].strip()
        if not nxt or is_special(nxt):
            break
        buf.append(nxt)
        i += 1
    para(" ".join(buf), size=BODY_SIZE, after=5)


# Párrafo final obligatorio (Word)
if not doc.paragraphs or doc.paragraphs[-1].text.strip():
    doc.add_paragraph()

doc.save(OUT)
print(f"OK · {OUT.name} · {OUT.stat().st_size/1024:.0f} KB")
print(f"párrafos={len(doc.paragraphs)} tablas={len(doc.tables)} secciones={len(doc.sections)}")
